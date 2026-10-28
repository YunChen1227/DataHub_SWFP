# -*- coding: utf-8 -*-
import sys
from pathlib import Path

sys.stdout.reconfigure(encoding="utf-8")

from openpyxl import load_workbook

paths = [
    Path(r"c:\Users\XMMXYY\Desktop\有票税号_标记是否在库.xlsx"),
    Path(r"c:\Users\XMMXYY\Desktop\有票税号(1)(1).xlsx"),
    Path(r"c:\Users\XMMXYY\Documents\xwechat_files\wxid_t7vgbilml8m511_c2b2\msg\file\2026-10\水滴存疑税号对账.xlsx"),
    Path(r"c:\Users\XMMXYY\Documents\xwechat_files\wxid_t7vgbilml8m511_c2b2\msg\file\2026-10\惠众9月账单-全部295个税号-近12个月复查.xlsx"),
    Path(r"c:\workspace\DataHub_SWFP\docs\税票分析接口文档.xlsx"),
]

for p in paths:
    print("=" * 80)
    print(p, "exists", p.exists(), "size", p.stat().st_size if p.exists() else 0)
    if not p.exists():
        continue
    try:
        wb = load_workbook(p, read_only=True, data_only=True)
    except Exception as e:
        print("OPEN FAIL", type(e).__name__, e)
        continue
    print("sheets:", wb.sheetnames)
    for sn in wb.sheetnames:
        sh = wb[sn]
        print("--- sheet", sn, "dims", sh.dimensions)
        n = 0
        for i, row in enumerate(sh.iter_rows(values_only=True), 1):
            print(i, row)
            n = i
            if i >= 8:
                break
        # count
        total = n
        if n >= 8:
            for _ in sh.iter_rows(min_row=9, values_only=True):
                total += 1
        print("rows_seen_approx", total)
    wb.close()
