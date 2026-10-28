# -*- coding: utf-8 -*-
"""Mark Excel cells that hit the 32767 limit and point at the full JSON."""
import sys
from pathlib import Path

from openpyxl import load_workbook

sys.stdout.reconfigure(encoding="utf-8")

xlsx = Path(r"c:\workspace\DataHub_SWFP\test_res\北京联信-各上游最新查得-接口返回.xlsx")
raw_dir = Path(r"c:\workspace\DataHub_SWFP\test_res\北京联信-各上游最新查得-raw")
wb = load_workbook(xlsx)
ws = wb.active
headers = [c.value for c in ws[1]]
col = headers.index("接口返回") + 1
file_col = len(headers) + 1
ws.cell(1, file_col, "完整返回文件")
files = sorted(p.name for p in raw_dir.glob("*.json"))
for i, name in enumerate(files, start=2):
    cell = ws.cell(i, col)
    text = cell.value or ""
    full = (raw_dir / name).read_text(encoding="utf-8")
    rel = "北京联信-各上游最新查得-raw/" + name
    ws.cell(i, file_col, rel)
    if len(full) > 32000:
        cell.value = "[Excel单元格上限32767，此处为截断。完整返回见「完整返回文件」列，或同目录 CSV 的「接口返回」列]\n" + text[:30000]
wb.save(xlsx)
print("updated", xlsx)
